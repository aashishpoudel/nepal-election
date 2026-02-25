#!/usr/bin/env python3
# -*- coding: utf-8 -*-

from __future__ import annotations

from nepal_election_constants import *  # provides NEPALI_DIGITS, etc.
import argparse
import re
import unicodedata
from typing import Iterable, Optional

import pandas as pd

# ✅ REQUIRED for bold styling
from openpyxl import load_workbook
from openpyxl.styles import Font


SOURCE_COL = "शैक्षिक योग्यता (माथिल्लो शैक्षिक योग्यता)"
TARGET_COL = "शैक्षिक योग्यता समूह"


def normalize_text(s: Optional[str]) -> str:
    if s is None:
        return ""
    s = str(s)
    s = unicodedata.normalize("NFC", s).strip().lower()

    # Remove dot and spaces
    s = s.replace(".", "")
    s = re.sub(r"\s+", "", s)

    # Replace ब -> व
    s = s.replace("ब", "व")

    # Convert Nepali numerals to English
    s = s.translate(NEPALI_DIGITS)

    # Convert Nepali number words to digits
    for k in sorted(NEPALI_NUMBER_WORDS.keys(), key=len, reverse=True):
        s = s.replace(k, NEPALI_NUMBER_WORDS[k])

    # दीर्घ -> ह्रस्व (approx)
    for k, v in DIRGHA_TO_HRASWA.items():
        s = s.replace(k, v)

    # Extra normalizations
    for a, b in EXTRA_NORMALIZATIONS:
        s = s.replace(a, b)

    return s


def normalize_keywords(keywords: Iterable[str]) -> list[str]:
    out: list[str] = []
    for kw in keywords:
        nk = normalize_text(kw)
        if nk:
            out.append(nk)
    seen = set()
    uniq = []
    for x in out:
        if x not in seen:
            uniq.append(x)
            seen.add(x)
    return uniq


def contains_any(haystack: str, needles: Iterable[str]) -> bool:
    return any(n in haystack for n in needles)


def classify_education_group(education_raw: Optional[str]) -> str:
    edu = normalize_text(education_raw)

    # ✅ FIX #1: NA -> Not Available
    if not edu:
        return "Not Available"

    phd_kw = [
        "पिएचडि", "विध्यावारिधि", "विधावारिधि", "विद्यावारिधि", "phd", "विदयावारिणि", "पिचडी"
    ]

    masters_kw = [
        "स्‍नात्तकोत्तर", "स्‍नाक्तोर", "स्‍नाकोत्तर", "स्नोकोत्तर", "स्नात्तकोत्तर",
        "स्‍नातोकत्तर", "स्‍नातकोत्तर", "स्‍नातकोतर", "स्नात्तकोतर", "स्नात्कोत्तर",
        "स्नातोकोत्तर", "स्नातोकत्तर", "स्नातकोत्तरोपाधि", "स्नातकोतर", "स्नातकोक्तर",
        "स्नातकोउत्तर", "स्नातकत्तोर", "स्नाकोत्तर", "स्ननातोकोत्तर", "स्नताकोत्तर",
        "स्तानकोत्तर", "ma", "आचार्य", "md",
        "मास्तर", "मास्टर्स", "मास्टर", "माष्टर्स", "माष्टर", "मस्टर",
        "डिग्री", "एमए", "mcom", "mfil", "mphil", "mtech", "mteck", "mtecknology",
        "एम‍ए", "एम्फिल", "एमफिल", "एमवियस", "एमबिइ", "एमएड", "एमविविएस",
        "mbbs", "ma", "med", "ms", "एमविए", "एमपिए", "एमटेक",
        "mpa", "llm", "एलएलवी", "llb", "mca", "mbs", "एमई", "dm"
    ]

    bachelors_kw = [
        "bbs", "bba",
        "स्‍नात्तक", "स्‍नातक", "स्नातक", "स्नताक", "स्तानतक", "स्तानक",
        "इन्जिनियरिङ", "ईन्जिनियर", "ईन्जीनियर", "सनातक", "ईन्जिनेरिङ",
        "व्याचुलर", "व्याचलर", "वैचलर",
        "विविएस", "विएड", "विसएसि", "विएस्सि", "विएससि", "विएस्सी", "विविए",
        "ba", "विकम", "विएससी", "विएल", "विए", "विइ",
        "विफारमेसी", "विएएलएलबी", "विएललवि", "विफार्म",
        "bacherlor", "bcahelor", "bsc", "engineer",
        "btu", "ballb", "bed", "btech", "bcom", "bl", "atpl", "be", "doctor", "bds"
    ]

    intermediate_kw = [
        "+2", "+12", "आइएड", "isc", "ia", "icom", "१२पास",
        "१२कक्षा", "१२उतिर्ण", "१२", "१०+२",
        "प्लसटु", "डिप्लोमा", "आइकम", "आई कम", "इन्टरमिडिएट",
        "प्रमाणपत्र", "आइए", "आइऐड", "inter",
        "इन्टरमेडियट", "इन्टरमीडियट", "इन्टरमेडियेट", "इण्टरमेडियट",
        "प्रविन्ताप्रमाणतह", "प्रविन्ताप्रमाणपत्रतह", "pcl", "plustwo",
        "प्रविणाताप्रमाणपत्र", "प्रविणतापत्र", "प्रविणताप्रमाणपत्र",
        "प्रविणताप्रमाण", "प्रविणता", "प्रमानपत्र", "आए", "आइऐ", "proficiency",
        "नर्सिङ", "उच्चशिक्षा", "आयएससि", "hseb", "ied", "ha", "डिल्लोमा इन्जनियरिङ्ग"
    ]

    slc_kw = [
        "slc", "see", "११", "१०pass", "१०कक्षा", "१०", "म्याट्रिक", "एसइइ", "sklc",
        "भेटनरी,जेटिए", "प्रवेशिका", "एसएलसि", "दशौंकक्षा", "एसएससि", "एलएलसि", "ऐसएलसि",
        "अहेव", "एसलसि"
    ]

    lt10_kw = [
        "सामान्यलेखपढ", "सामान्यलेखपण", "सामान्यसाक्षर", "सााधारणलेखपेढ",
        "सामान्य", "साधारण", "साधरण", "सातpass", "साक्षार", "साक्षरता", "साक्षर",
        "समान्य", "सधारण", "शिक्षित", "simple",
        "९pass", "८pass", "७pass", "६pass", "५pass", "४pass", "३pass", "२pass", "१pass",
        "९कक्षा", "८कक्षा", "७कक्षा", "६कक्षा", "५कक्षा", "४कक्षा", "३कक्षा",
        "कक्षा९", "कक्षा८", "कक्षा७", "कक्षा६", "कक्षा५", "कक्षा४", "कक्षा३",
        "लेखपढ", "दिक्षित",
        "माध्यमिकतह", "माध्यमिकशिक्षा", "मावि", "माव",
        "प्राथमिकशिक्षा", "प्राथमिकपरिक्षाpass", "प्राथमिकतह", "प्रावि", "आधारभुत", "स्वअध्ययन"
    ]

    noedu_kw = ["निराक्षर", "निरक्षर", "नभएको", "छैन", "अनपढ"]

    phd_kw_n = normalize_keywords(phd_kw)
    masters_kw_n = normalize_keywords(masters_kw)
    bachelors_kw_n = normalize_keywords(bachelors_kw)
    intermediate_kw_n = normalize_keywords(intermediate_kw)
    slc_kw_n = normalize_keywords(slc_kw)
    lt10_kw_n = normalize_keywords(lt10_kw)
    noedu_kw_n = normalize_keywords(noedu_kw)

    if contains_any(edu, noedu_kw_n):
        return "No Education"
    if contains_any(edu, phd_kw_n):
        return "PhD"
    if contains_any(edu, masters_kw_n):
        return "Masters"
    if contains_any(edu, bachelors_kw_n):
        return "Bachelors"
    if contains_any(edu, intermediate_kw_n):
        return "Intermediate"
    if contains_any(edu, slc_kw_n):
        return "SLC"
    if contains_any(edu, lt10_kw_n):
        return "<10 class"

    return "Not Available"


def apply_excel_bold_style(excel_path: str, column_name: str, font_size: Optional[int] = None) -> None:
    wb = load_workbook(excel_path)
    ws = wb.active

    target_col_index = None
    for i, cell in enumerate(ws[1], start=1):
        if cell.value == column_name:
            target_col_index = i
            break

    if target_col_index is None:
        raise ValueError(f"Column {column_name!r} not found in the output Excel.")

    header_font = Font(bold=True, size=(font_size or 12))
    ws.cell(row=1, column=target_col_index).font = header_font

    data_font = Font(bold=True, size=font_size) if font_size else Font(bold=True)

    for row in ws.iter_rows(min_row=2, min_col=target_col_index, max_col=target_col_index):
        for cell in row:
            if cell.value is not None and str(cell.value).strip() != "":
                cell.font = data_font

    wb.save(excel_path)


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--input", required=True, help="Input Excel file (.xlsx)")
    ap.add_argument("--output", required=True, help="Output Excel file (.xlsx)")
    ap.add_argument("--font-size", type=int, default=None)
    args = ap.parse_args()

    df = pd.read_excel(args.input, engine="openpyxl")

    if SOURCE_COL not in df.columns:
        raise KeyError(f"Missing required column: {SOURCE_COL!r}. Found: {list(df.columns)}")

    df[TARGET_COL] = df[SOURCE_COL].apply(classify_education_group)

    df.to_excel(args.output, index=False, engine="openpyxl")

    # ✅ FIX #2: Apply bold style to the OUTPUT file
    apply_excel_bold_style(args.output, TARGET_COL, font_size=args.font_size)

    print(df[TARGET_COL].value_counts(dropna=False))
    print(f"Wrote: {args.output}")


if __name__ == "__main__":
    main()